"""Leitura dos chamados e gravação dos resultados na planilha Excel (.xlsx)."""
from __future__ import annotations

import re
import time
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from chamado import normalizar_texto

# Resultados gravados na coluna "Resultado automação"
SUCESSO = "SUCESSO"
REVISAO_MANUAL = "REVISÃO MANUAL"
PENDENTE_GLPI = "PENDENTE GLPI"  # senha já trocada no CASI; falta (ou falhou) responder o chamado
ERRO = "ERRO"
IGNORADO = "IGNORADO"
PREFIXO_SIMULACAO = "SIMULADO - "
FINALIZADOS = {SUCESSO, REVISAO_MANUAL}  # linhas com esses resultados não são processadas de novo
# PENDENTE_GLPI e ERRO são reprocessados; quem precisa de um humano:
REVISAO_HUMANA = {REVISAO_MANUAL, ERRO}

# Nomes de cabeçalho aceitos (comparação sem acento, maiúsculas ou pontuação)
_ENTRADA = {
    "chamado": ("chamado", "id", "ticket", "n chamado", "no chamado", "numero do chamado", "numero"),
    "categoria": ("categoria",),
    "nome": ("nome", "nome completo", "colaborador", "nome do colaborador"),
    "matricula": ("matricula", "matricula do colaborador"),
    "descricao": ("descricao", "conteudo", "descricao do chamado"),
    # Opcional: força o(s) caractere(s) da senha para este chamado (ex.: "#" quando Dpsp@ não funciona)
    "caractere": ("caractere", "caractere senha", "caractere da senha", "caracter", "caracter senha", "caracteres"),
}
# Colunas de saída: nomes próprios para não confundir com a coluna "Status" do GLPI
_SAIDA = {
    "resultado": "Resultado automação",
    "detalhe": "Detalhe automação",
    "processado_em": "Processado em",
    "mensagem": "Mensagem GLPI",
}
_MODELO = ["Chamado", "Categoria", "Nome", "Matrícula", "Caractere senha", "Descrição", *_SAIDA.values()]
_CORES = {
    SUCESSO: "C6EFCE",
    REVISAO_MANUAL: "FFEB9C",
    PENDENTE_GLPI: "FCE4D6",
    ERRO: "FFC7CE",
    IGNORADO: "D9D9D9",
}


class PlanilhaErro(Exception):
    pass


@dataclass(frozen=True)
class LinhaChamado:
    numero: int  # número da linha no Excel
    chamado: str
    categoria: str
    nome: str
    matricula: str
    descricao: str
    caractere: str
    resultado: str
    mensagem: str  # mensagem já gerada (usada para reenviar ao GLPI sem mexer de novo no CASI)

    @property
    def vazia(self) -> bool:
        return not (self.nome or self.matricula or self.descricao)

    @property
    def referencia(self) -> str:
        return f"Linha {self.numero}" + (f" (chamado {self.chamado})" if self.chamado else "")


def normalizar_cabecalho(valor) -> str:
    """'Nº Chamado:' -> 'n chamado' (sem acento, pontuação e espaços extras)."""
    return " ".join(re.sub(r"[^a-z0-9 ]", " ", normalizar_texto(str(valor or ""))).split())


def _texto(valor) -> str:
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))  # Excel guarda matrícula numérica como 126594.0
    if isinstance(valor, datetime):
        return valor.isoformat(sep=" ", timespec="minutes")
    return str(valor).strip()


def criar_modelo(caminho: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Chamados"
    ws.append(_MODELO)
    for celula in ws[1]:
        celula.font = Font(bold=True)
    larguras = {"A": 14, "B": 40, "C": 38, "D": 14, "E": 16, "F": 50, "G": 22, "H": 70, "I": 18, "J": 70}
    for coluna, largura in larguras.items():
        ws.column_dimensions[coluna].width = largura
    ws.freeze_panes = "A2"
    wb.save(caminho)


class Planilha:
    def __init__(self, caminho: Path):
        self.caminho = caminho
        try:
            self.wb = load_workbook(caminho)
        except PermissionError as exc:
            raise PlanilhaErro(f"Não consegui abrir {caminho.name}. Feche a planilha no Excel e rode de novo.") from exc
        self.ws = self.wb.active
        self.linha_cabecalho, self.colunas = self._mapear_colunas()
        self.salvar()  # já testa se dá para gravar antes de mexer em qualquer coisa no CASI

    def _mapear_colunas(self) -> tuple[int, dict[str, int]]:
        for numero_linha in range(1, min(self.ws.max_row, 10) + 1):
            encontrados: dict[str, int] = {}
            for celula in self.ws[numero_linha]:
                titulo = normalizar_cabecalho(celula.value)
                for campo, nomes in _ENTRADA.items():
                    if campo not in encontrados and titulo in nomes:
                        encontrados[campo] = celula.column
                for campo, nome in _SAIDA.items():
                    if campo not in encontrados and titulo == normalizar_cabecalho(nome):
                        encontrados[campo] = celula.column
            if "matricula" in encontrados or "descricao" in encontrados:
                proxima = self.ws.max_column + 1
                for campo, nome in _SAIDA.items():  # cria as colunas de resultado se não existirem
                    if campo not in encontrados:
                        self.ws.cell(row=numero_linha, column=proxima, value=nome).font = Font(bold=True)
                        encontrados[campo] = proxima
                        proxima += 1
                return numero_linha, encontrados
        raise PlanilhaErro(
            "Não encontrei o cabeçalho na planilha. A primeira linha precisa ter pelo menos a coluna "
            "'Matrícula' (ou 'Descrição' com o texto do chamado). Colunas aceitas: " + ", ".join(_MODELO[:5])
        )

    def _valor(self, linha: int, campo: str) -> str:
        coluna = self.colunas.get(campo)
        return _texto(self.ws.cell(row=linha, column=coluna).value) if coluna else ""

    def linhas(self) -> list[LinhaChamado]:
        return [
            LinhaChamado(
                numero=n,
                chamado=self._valor(n, "chamado"),
                categoria=self._valor(n, "categoria"),
                nome=self._valor(n, "nome"),
                matricula=self._valor(n, "matricula"),
                descricao=self._valor(n, "descricao"),
                caractere=self._valor(n, "caractere"),
                resultado=self._valor(n, "resultado").upper(),
                mensagem=self._valor(n, "mensagem"),
            )
            for n in range(self.linha_cabecalho + 1, self.ws.max_row + 1)
        ]

    @property
    def tem_coluna_categoria(self) -> bool:
        return "categoria" in self.colunas

    @property
    def tem_coluna_chamado(self) -> bool:
        return "chamado" in self.colunas

    def chamados_existentes(self) -> set[str]:
        """Números de chamado já presentes na planilha (só dígitos)."""
        return {re.sub(r"\D", "", l.chamado) for l in self.linhas() if re.sub(r"\D", "", l.chamado)}

    def _ultima_linha_preenchida(self) -> int:
        colunas = list(self.colunas.values())
        for n in range(self.ws.max_row, self.linha_cabecalho, -1):
            if any(self.ws.cell(row=n, column=c).value not in (None, "") for c in colunas):
                return n
        return self.linha_cabecalho

    def adicionar(self, chamado: str, categoria: str, nome: str, matricula: str, descricao: str) -> None:
        """Acrescenta um chamado no fim da planilha. Chame salvar() depois."""
        linha = self._ultima_linha_preenchida() + 1
        valores = {
            "chamado": int(chamado) if chamado.isdigit() else chamado,
            "categoria": categoria,
            "nome": nome,
            "matricula": matricula,  # texto, para não perder zero à esquerda
            "descricao": descricao,
        }
        for campo, valor in valores.items():
            coluna = self.colunas.get(campo)
            if coluna and valor not in ("", None):
                celula = self.ws.cell(row=linha, column=coluna, value=valor)
                if campo == "matricula":
                    celula.number_format = "@"
                if campo == "descricao":
                    celula.alignment = Alignment(wrap_text=False, vertical="top")

    def registrar(self, linha: LinhaChamado, resultado: str, detalhe: str, mensagem: str, simulacao: bool) -> None:
        texto = f"{PREFIXO_SIMULACAO}{resultado}" if simulacao else resultado
        celula = self.ws.cell(row=linha.numero, column=self.colunas["resultado"], value=texto)
        cor = "DDEBF7" if simulacao else _CORES.get(resultado)
        if cor:
            celula.fill = PatternFill("solid", start_color=cor, end_color=cor)
        self.ws.cell(row=linha.numero, column=self.colunas["detalhe"], value=detalhe)
        self.ws.cell(row=linha.numero, column=self.colunas["processado_em"],
                     value=datetime.now().strftime("%d/%m/%Y %H:%M:%S"))
        # Sempre sobrescreve, para não sobrar a mensagem de uma execução anterior
        celula_msg = self.ws.cell(row=linha.numero, column=self.colunas["mensagem"], value=mensagem or None)
        celula_msg.alignment = Alignment(wrap_text=True, vertical="top")
        self.salvar()

    def exportar_revisao(self, destino: Path) -> int:
        """Gera uma planilha só com os chamados que precisam de um humano (REVISÃO MANUAL e ERRO).
        Devolve quantos foram exportados. Se não houver nenhum, apaga um arquivo antigo e não cria outro."""
        prefixo = PREFIXO_SIMULACAO.strip().upper()  # ignora o "SIMULADO -" na frente do status

        def status(l: LinhaChamado) -> str:
            return l.resultado.removeprefix(prefixo).strip()

        pendentes = [l for l in self.linhas() if status(l) in REVISAO_HUMANA]
        if not pendentes:
            destino.unlink(missing_ok=True)
            return 0
        wb = Workbook()
        ws = wb.active
        ws.title = "Para revisão"
        cabecalho = ["Chamado", "Categoria", "Nome", "Matrícula", "Resultado", "Motivo / o que fazer", "Descrição"]
        ws.append(cabecalho)
        for celula in ws[1]:
            celula.font = Font(bold=True)
        for l in pendentes:
            ws.append([int(l.chamado) if l.chamado.isdigit() else l.chamado, l.categoria, l.nome, l.matricula,
                       status(l), self._valor(l.numero, "detalhe"), l.descricao])
            ws.cell(row=ws.max_row, column=5).fill = PatternFill("solid", start_color=_CORES[status(l)],
                                                                 end_color=_CORES[status(l)])
        for coluna, largura in {"A": 12, "B": 38, "C": 34, "D": 12, "E": 16, "F": 70, "G": 60}.items():
            ws.column_dimensions[coluna].width = largura
        ws.freeze_panes = "A2"
        wb.save(destino)
        return len(pendentes)

    def salvar(self) -> None:
        for tentativa in range(5):  # Excel aberto ou OneDrive sincronizando seguram o arquivo
            try:
                self.wb.save(self.caminho)
                return
            except PermissionError as exc:
                if tentativa == 4:
                    raise PlanilhaErro(
                        f"Não consegui gravar em {self.caminho.name}: o arquivo está aberto no Excel? "
                        "Feche a planilha e rode de novo."
                    ) from exc
                time.sleep(2)
